# ============================================================
# SAGE — Agentic AI Teaching Kit Generator (Web App)
# Interactive Streamlit version: review & edit every stage
# (plan -> research -> teaching kit -> presentation -> assessment
# -> evaluation) before exporting the final files.
# ============================================================

import os
import re
import json
import time
import zipfile
import tempfile
from datetime import date
from typing import List, Optional, Literal

import requests
import faiss
import streamlit as st

from pydantic import BaseModel, Field
from openai import OpenAI
from sentence_transformers import SentenceTransformer

from pptx import Presentation as PptxPresentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn as _qn

from docx import Document
from docx.shared import Inches as DocxInches, Pt as DocxPt, RGBColor as DocxRGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ============================================================
# PAGE CONFIG
# ============================================================
st.set_page_config(
    page_title="SAGE — Agentic AI Teaching Kit Generator",
    page_icon="✦",
    layout="wide",
    initial_sidebar_state="expanded"
)

OPENAI_MODEL = "gpt-4o-mini"
EMBED_MODEL = "all-MiniLM-L6-v2"
HEADERS = {"User-Agent": "SAGE-Educational-Agent/1.0"}
USAGE_FILE = os.path.join(tempfile.gettempdir(), "sage_usage_log.json")


# ============================================================
# SECRETS / CONFIG
# ============================================================
def get_secret(name, default=None):
    try:
        if name in st.secrets:
            return st.secrets[name]
    except Exception:
        pass
    return os.environ.get(name, default)


OPENAI_API_KEY = get_secret("OPENAI_API_KEY")
APP_ACCESS_CODE = get_secret("APP_ACCESS_CODE", "")
MAX_DAILY_RUNS = int(get_secret("MAX_DAILY_RUNS", 20))
MAX_SESSION_RUNS = int(get_secret("MAX_SESSION_RUNS", 3))


if not OPENAI_API_KEY:
    st.error(
        "OpenAI API key not found. Please add OPENAI_API_KEY to your Secrets "
        "(Hugging Face Space settings or .streamlit/secrets.toml locally)."
    )
    st.stop()


@st.cache_resource(show_spinner=False)
def get_client():
    return OpenAI(api_key=OPENAI_API_KEY)


@st.cache_resource(show_spinner="Loading embedding model (one-time only)...")
def get_embedder():
    return SentenceTransformer(EMBED_MODEL)


client = get_client()


# ============================================================
# ACCESS GATE (simple passcode, optional)
# ============================================================
if APP_ACCESS_CODE:
    if not st.session_state.get("authenticated"):
        st.markdown(
            """
            <div style="max-width:440px;margin:6vh auto 0 auto;padding:2rem;
                        border:1px solid #E5E7EB;border-radius:16px;background:#FFFFFF;
                        box-shadow:0 4px 20px rgba(15,23,42,0.06);">
                <div style="font-size:1.35rem;font-weight:800;color:#0F172A;letter-spacing:0.5px;">
                    ✦ SAGE
                </div>
                <div style="color:#64748B;font-size:0.88rem;margin-top:4px;margin-bottom:1.4rem;">
                    Agentic AI Teaching Kit Generator
                </div>
                <div style="color:#334155;font-size:0.85rem;">
                    This site is protected by an access code. Please enter the code to continue.
                </div>
            </div>
            """,
            unsafe_allow_html=True
        )
        _, col, _ = st.columns([1, 2, 1])
        with col:
            code_input = st.text_input("Access code", type="password")
            if st.button("Sign in", type="primary", use_container_width=True):
                if code_input == APP_ACCESS_CODE:
                    st.session_state.authenticated = True
                    st.rerun()
                else:
                    st.error("Incorrect access code.")
        st.stop()


# ============================================================
# DAILY / SESSION USAGE CONTROL (basic cost safety net —
# always also set hard spending limits on the OpenAI dashboard)
# ============================================================
def daily_usage_ok_and_increment():
    today = date.today().isoformat()
    data = {}

    if os.path.exists(USAGE_FILE):
        try:
            with open(USAGE_FILE, "r") as f:
                data = json.load(f)
        except Exception:
            data = {}

    count_today = data.get(today, 0)

    if count_today >= MAX_DAILY_RUNS:
        return False

    with open(USAGE_FILE, "w") as f:
        json.dump({today: count_today + 1}, f)

    return True


# ============================================================
# HELPERS
# ============================================================
def sanitize_filename(topic: str) -> str:
    """Converts a prompt/topic string into a clean, filesystem-safe title string (supports Arabic)."""
    clean = re.sub(r'[^\w\s\u0600-\u06FF_-]', '', topic).strip()
    clean = re.sub(r'[\s]+', '_', clean)
    return clean.strip('_')[:40] if clean else "Topic"


def is_arabic(text: str) -> bool:
    """Checks if the given text contains Arabic characters."""
    return bool(re.search(r'[\u0600-\u06FF]', text or ""))


def normalize_arabic(text: str) -> str:
    """Normalizes Arabic characters (alef forms, diacritics) for robust keyword matching."""
    text = re.sub(r'[\u064B-\u0652]', '', text)
    text = re.sub(r'[أإآ]', 'ا', text)
    text = re.sub(r'ة', 'ه', text)
    return text.lower()


def list_to_text(items):
    return "\n".join(items)


def text_to_list(text):
    return [line.strip() for line in text.split("\n") if line.strip()]


# ============================================================
# THEME SYSTEM (5 PowerPoint themes + auto detection)
# ============================================================
THEMES = {
    "emerald": {
        "key": "emerald",
        "name": "Emerald",
        "label": "Biology / Medical / Health",
        "primary": "#059669",
        "dark": "#047857",
        "soft_bg": "#ECFDF5",
        "text": "#064E3B",
        "accent_light": "#A7F3D0",
        "muted": "#64748B",
    },
    "tech_blue": {
        "key": "tech_blue",
        "name": "Tech Blue",
        "label": "Computer / Tech / Engineering / AI",
        "primary": "#2563EB",
        "dark": "#1D4ED8",
        "soft_bg": "#EFF6FF",
        "text": "#1E3A8A",
        "accent_light": "#BFDBFE",
        "muted": "#64748B",
    },
    "indigo": {
        "key": "indigo",
        "name": "Indigo",
        "label": "Finance / Business / Economics",
        "primary": "#4F46E5",
        "dark": "#4338CA",
        "soft_bg": "#EEF2FF",
        "text": "#312E81",
        "accent_light": "#C7D2FE",
        "muted": "#64748B",
    },
    "amber": {
        "key": "amber",
        "name": "Amber",
        "label": "Arts / Humanities / History",
        "primary": "#D97706",
        "dark": "#B45309",
        "soft_bg": "#FFFBEB",
        "text": "#78350F",
        "accent_light": "#FDE68A",
        "muted": "#64748B",
    },
    "slate": {
        "key": "slate",
        "name": "Slate",
        "label": "General / Fallback",
        "primary": "#475569",
        "dark": "#334155",
        "soft_bg": "#F8FAFC",
        "text": "#1E293B",
        "accent_light": "#CBD5E1",
        "muted": "#64748B",
    },
}


THEME_KEYWORDS = {
    "emerald": [
        "biology", "bio", "medical", "medicine", "health", "anatomy", "physiology",
        "genetics", "gene", "dna", "microbiology", "nutrition", "clinical",
        "organism", "botany", "zoology", "virus", "bacterial", "pharma", "cell",
        "احياء", "طب", "صحه", "خليه", "خلايا", "جينات", "جين", "تشريح", "وراثه",
        "فيروس", "بكتيريا", "دواء", "صيدله", "حيوي", "احيائي", "علاجي", "جسم",
        "مرض", "امراض"
    ],
    "tech_blue": [
        "computer", "programming", "code", "software", "ai", "artificial intelligence",
        "machine learning", "deep learning", "data science", "data", "cyber",
        "cybersecurity", "algorithm", "python", "network", "web", "developer",
        "engineering", "robotics", "tech", "technology",
        "برمجه", "حاسوب", "كمبيوتر", "تقنيه", "ذكاء", "اصطناعي", "بيانات",
        "خوارزميه", "شبكه", "شبكات", "تطوير", "موقع", "برنامج", "سيبراني",
        "الكتروني", "هندسه"
    ],
    "indigo": [
        "finance", "business", "economics", "econ", "accounting", "investment",
        "banking", "marketing", "entrepreneurship", "market", "management",
        "trade", "corporate", "startup", "commerce",
        "تجاره", "اعمال", "اقتصاد", "تسويق", "ماليه", "استثمار", "اداره",
        "محاسبه", "شركة", "شركات"
    ],
    "amber": [
        "history", "art", "arts", "literature", "lit", "philosophy", "culture",
        "humanities", "heritage", "language", "music", "design", "poetry",
        "فن", "فنون", "تاريخ", "ادب", "فلسفه", "موسيقى", "تصميم", "ثقافه",
        "لغه", "شعر"
    ],
}


def detect_theme(topic: str) -> str:
    """
    Deterministically detect the best theme key from a topic string.
    Falls back to 'slate' (General) when no category matches.
    Checks categories in priority order: emerald, tech_blue, indigo, amber.
    """
    if not topic:
        return "slate"

    normalized = normalize_arabic(topic)

    for key in ("emerald", "tech_blue", "indigo", "amber"):
        for kw in THEME_KEYWORDS[key]:
            kw_norm = normalize_arabic(kw)
            if kw_norm and kw_norm in normalized:
                return key

    return "slate"


def get_theme_palette(theme_key: str) -> dict:
    """Return a copy of the palette dict for the given theme key."""
    return dict(THEMES.get(theme_key, THEMES["slate"]))


def _hex_to_rgb(hex_str: str) -> RGBColor:
    h = hex_str.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


# ============================================================
# DATA MODELS
# ============================================================
class Objective(BaseModel):
    objective: str


class Plan(BaseModel):
    topic: str
    audience: str
    duration: str
    objectives: List[Objective]
    key_concepts: List[str]
    lesson_flow: List[str]


class TeachingKit(BaseModel):
    summary: str
    key_terms: List[str]
    explanation: List[str]
    examples: List[str]
    teaching_tips: List[str]
    misconceptions: List[str]


class SlideCard(BaseModel):
    card_title: str = Field(description="Short 2-4 word bold title for a visual card container")
    card_body: str = Field(description="1-2 concise summary sentences")


class ReferenceItem(BaseModel):
    title: str
    author_or_source: str
    description: str


class Slide(BaseModel):
    title: str
    subtitle: str
    layout_type: Literal["cards", "bullets", "paragraph", "example", "references"]
    cards: Optional[List[SlideCard]] = None
    bullets: Optional[List[str]] = None
    paragraph_text: Optional[str] = None
    example_title: Optional[str] = None
    example_body: Optional[str] = None
    references: Optional[List[ReferenceItem]] = None
    key_takeaway: str
    speaker_notes: str


class SlideDeckSchema(BaseModel):
    topic: str
    slides: List[Slide]


class MultipleChoiceQuestion(BaseModel):
    question_text: str
    options: List[str]
    correct_answer: str
    explanation: str


class TrueFalseQuestion(BaseModel):
    statement_text: str
    correct_answer: Literal["True", "False"]
    explanation: str


class AssessmentSchema(BaseModel):
    title: str
    instructions: str
    mc_questions: List[MultipleChoiceQuestion]
    tf_questions: List[TrueFalseQuestion]


class Evaluation(BaseModel):
    score: int = Field(ge=0, le=100)
    approved: bool
    issues: List[str]
    improvements: List[str]


class RevisionPackage(BaseModel):
    teaching_kit: TeachingKit
    presentation: SlideDeckSchema
    assessment: AssessmentSchema


# ============================================================
# OPENAI STRUCTURED OUTPUT
# ============================================================
def llm(prompt, schema, retries=3):
    last_error = None

    for attempt in range(retries):
        try:
            response = client.responses.parse(
                model=OPENAI_MODEL,
                input=[
                    {
                        "role": "system",
                        "content": (
                            "You are SAGE, an educational AI agent and expert educational "
                            "content designer. Create accurate, clear, useful educational "
                            "content. Follow the requested structure exactly. "
                            "Use retrieved sources as factual grounding. "
                            "Do not invent unsupported facts."
                        )
                    },
                    {"role": "user", "content": prompt}
                ],
                text_format=schema
            )

            for output in response.output:
                if getattr(output, "type", None) == "message":
                    for content_item in output.content:
                        parsed = getattr(content_item, "parsed", None)
                        if parsed is not None:
                            return parsed

            raise RuntimeError("The model returned no valid structured output.")

        except Exception as e:
            last_error = e
            if attempt < retries - 1:
                time.sleep(2)

    raise RuntimeError(f"OpenAI request failed: {last_error}")


def llm_with_validation(prompt, schema, validate_fn, retries=3):
    last_error = None
    current_prompt = prompt

    for attempt in range(retries):
        result = llm(current_prompt, schema)
        try:
            validate_fn(result)
            return result
        except RuntimeError as e:
            last_error = e
            current_prompt = (
                prompt
                + "\n\nIMPORTANT — YOUR PREVIOUS ATTEMPT FAILED VALIDATION:\n"
                + f"{e}\n"
                + "Re-generate the FULL response from scratch, making sure "
                + "this exact requirement is satisfied this time."
            )

    raise RuntimeError(f"Validation kept failing after {retries} attempts: {last_error}")


# ============================================================
# PLANNER AGENT
# ============================================================
def planner_agent(topic):
    prompt = f"""
Create a teaching plan.

Topic: {topic}
Audience: University students
Duration: 60 minutes

Requirements:
- 4 measurable learning objectives
- 5 important key concepts
- 5-step lesson flow

Keep the plan clear and suitable for university students.
"""
    return llm(prompt, Plan)


# ============================================================
# RESEARCH TOOLS (Wikipedia + OpenAlex + FAISS RAG)
# ============================================================
def wikipedia_search(query, limit=3):
    try:
        response = requests.get(
            "https://en.wikipedia.org/w/api.php",
            params={
                "action": "query", "list": "search", "srsearch": query,
                "format": "json", "srlimit": limit
            },
            headers=HEADERS, timeout=20
        )
        response.raise_for_status()
        data = response.json()
        return [item["pageid"] for item in data.get("query", {}).get("search", [])]
    except Exception:
        return []


def wikipedia_content(page_id):
    try:
        response = requests.get(
            "https://en.wikipedia.org/w/api.php",
            params={
                "action": "query", "prop": "extracts", "explaintext": 1,
                "exchars": 10000, "pageids": page_id, "format": "json"
            },
            headers=HEADERS, timeout=20
        )
        response.raise_for_status()
        data = response.json()
        page = data["query"]["pages"][str(page_id)]
        return page.get("title", ""), page.get("extract", "")
    except Exception:
        return "", ""


def reconstruct_abstract(inverted_index):
    if not inverted_index:
        return ""
    words = []
    for word, positions in inverted_index.items():
        for position in positions:
            words.append((position, word))
    words.sort(key=lambda x: x[0])
    return " ".join(word for _, word in words)


def openalex_search(query, limit=4):
    results = []
    try:
        response = requests.get(
            "https://api.openalex.org/works",
            params={"search": query, "per-page": limit},
            headers=HEADERS, timeout=20
        )
        response.raise_for_status()
        works = response.json().get("results", [])
        for work in works:
            abstract = reconstruct_abstract(work.get("abstract_inverted_index"))
            if len(abstract.split()) < 30:
                continue
            results.append({
                "title": work.get("display_name", "Academic source"),
                "text": abstract,
                "url": work.get("doi") or work.get("id") or "",
                "type": "OpenAlex"
            })
    except Exception:
        pass
    return results


def chunk_text(text, size=450, overlap=70):
    words = text.split()
    chunks = []
    step = max(1, size - overlap)
    for start in range(0, len(words), step):
        chunk = " ".join(words[start:start + size])
        if len(chunk.split()) >= 40:
            chunks.append(chunk)
    return chunks


def research_agent(plan):
    embedder = get_embedder()

    queries = [plan.topic, *plan.key_concepts, *(item.objective for item in plan.objectives)]

    sources = []
    seen = set()

    for query in queries[:8]:
        page_ids = wikipedia_search(query, limit=2)
        for page_id in page_ids:
            title, text = wikipedia_content(page_id)
            if not title or not text:
                continue
            key = f"wiki:{title}"
            if key in seen:
                continue
            seen.add(key)
            for chunk in chunk_text(text):
                sources.append({
                    "title": title, "text": chunk,
                    "url": f"https://en.wikipedia.org/?curid={page_id}",
                    "type": "Wikipedia"
                })

        academic_sources = openalex_search(query, limit=2)
        for source in academic_sources:
            key = f"openalex:{source['title']}"
            if key in seen:
                continue
            seen.add(key)
            for chunk in chunk_text(source["text"]):
                sources.append({
                    "title": source["title"], "text": chunk,
                    "url": source["url"], "type": source["type"]
                })

    if not sources:
        raise RuntimeError("No sources were retrieved. Please check internet access.")

    texts = [s["text"] for s in sources]
    vectors = embedder.encode(texts, normalize_embeddings=True, show_progress_bar=False).astype("float32")

    index = faiss.IndexFlatIP(vectors.shape[1])
    index.add(vectors)

    query_text = (
        plan.topic + " " + " ".join(plan.key_concepts) + " "
        + " ".join(item.objective for item in plan.objectives)
    )
    query_vector = embedder.encode([query_text], normalize_embeddings=True, show_progress_bar=False).astype("float32")

    top_k = min(10, len(sources))
    scores, indices = index.search(query_vector, top_k)

    retrieved = []
    for index_id, score in zip(indices[0], scores[0]):
        if index_id < 0:
            continue
        source = dict(sources[index_id])
        source["score"] = float(score)
        retrieved.append(source)

    return retrieved


def build_context(sources):
    context = []
    for i, source in enumerate(sources, start=1):
        context.append(f"""
[SOURCE {i}]
Title: {source["title"]}
Type: {source["type"]}
URL: {source["url"]}

Content:
{source["text"]}
""")
    return "\n".join(context)


# ============================================================
# TEACHING KIT AGENT
# ============================================================
def teaching_kit_agent(plan, context):
    prompt = f"""
Create a complete teaching kit.

Topic: {plan.topic}
Audience: {plan.audience}
Duration: {plan.duration}

Learning objectives:
{chr(10).join("- " + item.objective for item in plan.objectives)}

Key concepts:
{chr(10).join("- " + item for item in plan.key_concepts)}

Retrieved research:
{context}

Requirements:
- clear summary
- key terms
- concept explanations
- practical examples
- teaching tips
- common misconceptions

Use the retrieved research for factual grounding.
"""
    return llm(prompt, TeachingKit)


# ============================================================
# PRESENTATION AGENT
# ============================================================
def presentation_agent(plan, teaching_kit, context, rtl_mode=False):
    prompt = f"""
You are an expert educational content designer. Create an in-depth,
visually structured presentation deck about the following topic.

Topic: {plan.topic}
Audience: {plan.audience}

Create EXACTLY 10 content slides (a references slide will be appended
separately from the real retrieved sources, so do NOT create one yourself).

Structure:
1. Introduction
2-8. Core concepts and explanations
9. Practical application / real-world example
10. Summary and key takeaways

Vary the layout_type across the deck — use a mix of "cards" (exactly 3
cards), "bullets" (3-5 concise bullets), "paragraph" (2-3 structured
sentences), and "example" (a code snippet or real-world walkthrough).
Avoid repeating the same layout on consecutive slides.

Every slide must include a short title, a one-sentence subtitle, a
single high-impact key_takeaway, and comprehensive speaker_notes.

CRITICAL: the "slides" list must contain EXACTLY 10 items. Count them
before responding.

Teaching summary:
{teaching_kit.summary}

Key terms:
{", ".join(teaching_kit.key_terms)}

Explanations:
{chr(10).join("- " + item for item in teaching_kit.explanation)}

Examples:
{chr(10).join("- " + item for item in teaching_kit.examples)}

Retrieved research (use for factual grounding, do not invent facts):
{context}
"""

    if rtl_mode:
        prompt += (
            "\n\nIMPORTANT: Write ALL slide content (titles, subtitles, cards, "
            "bullets, paragraphs, examples, key takeaways, speaker notes) in "
            "high-quality, natural Modern Standard Arabic."
        )

    def validate(result):
        if len(result.slides) != 10:
            raise RuntimeError(
                f"The 'slides' list must contain exactly 10 items, but it contained {len(result.slides)}."
            )

    try:
        return llm_with_validation(prompt, SlideDeckSchema, validate, retries=3)
    except RuntimeError:
        result = llm(prompt, SlideDeckSchema)
        if len(result.slides) > 10:
            result.slides = result.slides[:10]
        while len(result.slides) < 10:
            if rtl_mode:
                result.slides.append(Slide(
                    title="ملخص", subtitle="ملخص سريع للموضوع",
                    layout_type="paragraph", paragraph_text=teaching_kit.summary,
                    key_takeaway="راجع الملخص أعلاه قبل المتابعة.",
                    speaker_notes="استخدم هذه الشريحة لتلخيص الموضوع."
                ))
            else:
                result.slides.append(Slide(
                    title="Recap", subtitle="Quick recap of the topic",
                    layout_type="paragraph", paragraph_text=teaching_kit.summary,
                    key_takeaway="Review the summary above before moving on.",
                    speaker_notes="Use this slide to recap the topic summary."
                ))
        return result


def build_references_slide(sources, rtl_mode=False):
    seen_urls = set()
    reference_items = []
    for source in sources:
        url = source.get("url", "")
        if not url or url in seen_urls:
            continue
        seen_urls.add(url)
        reference_items.append(ReferenceItem(
            title=source.get("title", "Source"),
            author_or_source=source.get("type", "Source"),
            description=url
        ))
        if len(reference_items) == 4:
            break

    if rtl_mode:
        return Slide(
            title="المراجع", subtitle="المصادر التي تم استرجاعها واستُخدمت لبناء هذه الحقيبة",
            layout_type="references", references=reference_items,
            key_takeaway="كل معلومة في هذا العرض مبنية على مصدر حقيقي تم استرجاعه فعلياً.",
            speaker_notes="اختم الجلسة بتوجيه المتعلمين لهذه المصادر لمزيد من القراءة."
        )

    return Slide(
        title="References", subtitle="Sources retrieved and used to ground this teaching kit",
        layout_type="references", references=reference_items,
        key_takeaway="Every claim in this deck traces back to a real, retrieved source.",
        speaker_notes="Close the session by pointing learners to these sources for further reading."
    )


# ============================================================
# ASSESSMENT AGENT
# ============================================================
def assessment_agent(plan, teaching_kit, context, rtl_mode=False):
    prompt = f"""
You are an educational assessment expert. Build a comprehensive
assessment containing EXACTLY 20 questions in total:

1. EXACTLY 10 Multiple-Choice Questions (4 options each; the
   correct_answer must exactly match one of the options)
2. EXACTLY 10 True/False Questions

Topic: {plan.topic}
Audience: {plan.audience}

Learning objectives:
{chr(10).join("- " + item.objective for item in plan.objectives)}

Key concepts: {", ".join(plan.key_concepts)}

Teaching content: {teaching_kit.summary}

Retrieved research (use for factual grounding):
{context}

Requirements:
- exactly 10 multiple-choice questions, each with exactly 4 options
- exactly 10 true/false questions
- mix easy, medium, and challenging questions
- include a clear explanation for every question
- give the assessment a short title and brief instructions

CRITICAL: "mc_questions" must contain EXACTLY 10 items and "tf_questions"
must contain EXACTLY 10 items. Count them before responding.
"""

    if rtl_mode:
        prompt += (
            "\n\nIMPORTANT: Write ALL assessment questions, options, "
            "explanations, the title, and the instructions in natural, "
            "high-quality Modern Standard Arabic."
        )

    def validate(result):
        if len(result.mc_questions) != 10:
            raise RuntimeError(f"'mc_questions' must contain exactly 10 items, but it contained {len(result.mc_questions)}.")
        if len(result.tf_questions) != 10:
            raise RuntimeError(f"'tf_questions' must contain exactly 10 items, but it contained {len(result.tf_questions)}.")
        for question in result.mc_questions:
            if len(question.options) != 4:
                raise RuntimeError("Each multiple-choice question must have exactly 4 options.")
            if question.correct_answer not in question.options:
                raise RuntimeError("A multiple-choice question contains an invalid correct_answer.")

    try:
        return llm_with_validation(prompt, AssessmentSchema, validate, retries=3)
    except RuntimeError:
        result = llm(prompt, AssessmentSchema)
        if len(result.mc_questions) > 10:
            result.mc_questions = result.mc_questions[:10]
        if len(result.tf_questions) > 10:
            result.tf_questions = result.tf_questions[:10]
        validate(result)
        return result


# ============================================================
# EVALUATOR + REVISION AGENTS
# ============================================================
def evaluator_agent(plan, teaching_kit, presentation, assessment, context):
    prompt = f"""
Evaluate this educational package.

Topic: {plan.topic}

Evaluate: factual accuracy, grounding in retrieved research,
learning-objective alignment, educational clarity, presentation quality,
assessment quality, consistency, usefulness for university students.

Learning objectives:
{chr(10).join("- " + item.objective for item in plan.objectives)}

Teaching kit: {teaching_kit.model_dump_json()}
Presentation: {presentation.model_dump_json()}
Assessment: {assessment.model_dump_json()}
Research: {context}

Rules:
- score from 0 to 100
- approved must be true only when score >= 85
- identify concrete issues
- identify concrete improvements
"""
    return llm(prompt, Evaluation)


def revision_agent(plan, teaching_kit, presentation, assessment, evaluation, context):
    prompt = f"""
Revise the educational package based on evaluator feedback.

Topic: {plan.topic}
Evaluator feedback: {evaluation.model_dump_json()}
Current teaching kit: {teaching_kit.model_dump_json()}
Current presentation (exactly 10 content slides, no references slide): {presentation.model_dump_json()}
Current assessment: {assessment.model_dump_json()}
Research: {context}

Fix the identified problems while preserving accurate content.
Keep exactly 10 content slides and exactly 20 assessment questions
(10 multiple-choice + 10 true/false).
"""
    return llm(prompt, RevisionPackage)


# ============================================================
# PPTX LOW-LEVEL HELPERS
# ============================================================
DEFAULT_AR_FONT = "Arial"
DEFAULT_EN_FONT = "Calibri"


def _rtl_para(paragraph):
    """Mark a paragraph as RTL for proper Arabic rendering."""
    try:
        pPr = paragraph._p.get_or_add_pPr()
        pPr.set("rtl", "1")
    except Exception:
        pass


def _style_run(run, size=None, bold=None, italic=None, color=None, font=None):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    if font:
        run.font.name = font
        # try to set complex script font too (Arabic etc.)
        try:
            rPr = run._r.get_or_add_rPr()
            cs = rPr.find(_qn("a:cs"))
            if cs is None:
                cs = rPr.makeelement(_qn("a:cs"), {"typeface": font})
                rPr.append(cs)
            else:
                cs.set("typeface", font)
        except Exception:
            pass


def _add_textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))


def _add_shape(slide, shape_type, left, top, width, height, fill_rgb=None,
               line_rgb=None, line_width=1.0):
    shp = slide.shapes.add_shape(
        shape_type, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    if fill_rgb is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill_rgb
    if line_rgb is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_rgb
        shp.line.width = Pt(line_width)
    # remove default text
    shp.text_frame.text = ""
    return shp


def _set_bg(slide, rgb):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb


def _slide_header(slide, title, subtitle, theme_rgb, rtl):
    """Add a title + subtitle header consistent across all content slides."""
    align = PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT
    font = DEFAULT_AR_FONT if rtl else DEFAULT_EN_FONT

    # top accent bar
    _add_shape(
        slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.14,
        fill_rgb=theme_rgb["primary"]
    )
    # left brand mark (small square)
    brand_left = 12.55 if rtl else 0.7
    _add_shape(
        slide, MSO_SHAPE.ROUNDED_RECTANGLE,
        brand_left, 0.42, 0.16, 0.16,
        fill_rgb=theme_rgb["primary"]
    )

    # title
    title_box = _add_textbox(slide, 0.7, 0.75, 11.933, 0.9)
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    if rtl:
        _rtl_para(p)
    run = p.add_run()
    run.text = title
    _style_run(run, size=28, bold=True, color=theme_rgb["text"], font=font)

    # subtitle
    sub_box = _add_textbox(slide, 0.7, 1.55, 11.933, 0.55)
    stf = sub_box.text_frame
    stf.word_wrap = True
    sp = stf.paragraphs[0]
    sp.alignment = align
    if rtl:
        _rtl_para(sp)
    srun = sp.add_run()
    srun.text = subtitle
    _style_run(srun, size=14, bold=False, color=theme_rgb["muted"], font=font)

    return {"align": align, "font": font}


def _slide_takeaway(slide, takeaway_text, theme_rgb, rtl, top=6.05, height=0.85):
    """Bottom callout with the key takeaway."""
    align = PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT
    font = DEFAULT_AR_FONT if rtl else DEFAULT_EN_FONT

    _add_shape(
        slide, MSO_SHAPE.ROUNDED_RECTANGLE,
        0.7, top, 11.933, height,
        fill_rgb=theme_rgb["soft_bg"]
    )

    tb = _add_textbox(slide, 0.9, top + 0.12, 11.533, height - 0.2)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    if rtl:
        _rtl_para(p)

    label = "KEY TAKEAWAY: " if not rtl else "الخلاصة الأساسية: "
    r1 = p.add_run()
    r1.text = label
    _style_run(r1, size=11.5, bold=True, color=theme_rgb["primary"], font=font)

    r2 = p.add_run()
    r2.text = takeaway_text
    _style_run(r2, size=11.5, bold=False, color=theme_rgb["text"], font=font)


# ============================================================
# PPTX SLIDE LAYOUTS
# ============================================================
def _add_title_slide(prs, layout, deck_topic, theme_rgb, rtl):
    """Layout 1 — Title / Hero slide with strong branding."""
    slide = prs.slides.add_slide(layout)
    _set_bg(slide, theme_rgb["soft_bg"])

    # Left vertical accent block
    _add_shape(
        slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.35, 7.5,
        fill_rgb=theme_rgb["primary"]
    )
    # Bottom accent bar
    _add_shape(
        slide, MSO_SHAPE.RECTANGLE, 0, 7.25, 13.333, 0.25,
        fill_rgb=theme_rgb["dark"]
    )

    align = PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT
    font = DEFAULT_AR_FONT if rtl else DEFAULT_EN_FONT

    # SAGE brand mark
    brand_box = _add_textbox(slide, 1.0, 0.9, 6.0, 0.5)
    btf = brand_box.text_frame
    bp = btf.paragraphs[0]
    bp.alignment = align
    if rtl:
        _rtl_para(bp)
    brun = bp.add_run()
    brun.text = "✦ SAGE"
    _style_run(brun, size=16, bold=True, color=theme_rgb["primary"], font=font)

    # Eyebrow
    eyebrow_box = _add_textbox(slide, 1.0, 2.1, 11.3, 0.4)
    etf = eyebrow_box.text_frame
    ep = etf.paragraphs[0]
    ep.alignment = align
    if rtl:
        _rtl_para(ep)
    erun = ep.add_run()
    erun.text = "TEACHING KIT" if not rtl else "حقيبة تعليمية"
    _style_run(erun, size=12, bold=True, color=theme_rgb["dark"], font=font)

    # Main title
    title_box = _add_textbox(slide, 1.0, 2.55, 11.3, 1.9)
    ttf = title_box.text_frame
    ttf.word_wrap = True
    tp = ttf.paragraphs[0]
    tp.alignment = align
    if rtl:
        _rtl_para(tp)
    trun = tp.add_run()
    trun.text = deck_topic
    _style_run(trun, size=40, bold=True, color=theme_rgb["text"], font=font)

    # Subtitle
    sub_box = _add_textbox(slide, 1.0, 4.75, 11.3, 0.7)
    stf = sub_box.text_frame
    stf.word_wrap = True
    sp = stf.paragraphs[0]
    sp.alignment = align
    if rtl:
        _rtl_para(sp)
    srun = sp.add_run()
    srun.text = (
        "A complete teaching package generated by SAGE"
        if not rtl else
        "حقيبة تعليمية متكاملة تم توليدها بواسطة SAGE"
    )
    _style_run(srun, size=15, bold=False, color=theme_rgb["muted"], font=font)

    # Accent divider
    div_left = 11.0 if rtl else 1.0
    _add_shape(
        slide, MSO_SHAPE.RECTANGLE, div_left, 4.55, 1.35, 0.06,
        fill_rgb=theme_rgb["primary"]
    )


def _add_cards_slide(prs, layout, slide_data, theme_rgb, rtl):
    """Layout 3 — 3 cards side by side, or 2 cards if fewer items."""
    slide = prs.slides.add_slide(layout)
    _set_bg(slide, _hex_to_rgb("#FFFFFF"))

    hdr = _slide_header(slide, slide_data.title, slide_data.subtitle, theme_rgb, rtl)

    cards = slide_data.cards or []
    n = min(len(cards), 3)
    if n == 0:
        _slide_takeaway(slide, slide_data.key_takeaway, theme_rgb, rtl)
        _add_notes(slide, slide_data.speaker_notes)
        return

    total_w = 11.933
    gap = 0.35
    card_w = (total_w - gap * (n - 1)) / n
    top = 2.25
    card_h = 3.4

    for idx, card in enumerate(cards[:n]):
        visual_idx = (n - 1 - idx) if rtl else idx
        left = 0.7 + visual_idx * (card_w + gap)

        # card body
        _add_shape(
            slide, MSO_SHAPE.ROUNDED_RECTANGLE,
            left, top, card_w, card_h,
            fill_rgb=_hex_to_rgb("#FFFFFF"),
            line_rgb=theme_rgb["accent_light"],
            line_width=1.25
        )

        # number badge
        badge_left = left + card_w - 0.85 if rtl else left + 0.25
        _add_shape(
            slide, MSO_SHAPE.ROUNDED_RECTANGLE,
            badge_left, top + 0.22, 0.6, 0.42,
            fill_rgb=theme_rgb["soft_bg"]
        )
        bbox = _add_textbox(slide, badge_left, top + 0.24, 0.6, 0.4)
        bp = bbox.text_frame.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        brun = bp.add_run()
        brun.text = f"0{idx + 1}"
        _style_run(brun, size=13, bold=True, color=theme_rgb["primary"],
                   font=hdr["font"])

        # card title
        ctitle_box = _add_textbox(slide, left + 0.22, top + 0.85, card_w - 0.44, 0.9)
        ctf = ctitle_box.text_frame
        ctf.word_wrap = True
        cp1 = ctf.paragraphs[0]
        cp1.alignment = hdr["align"]
        if rtl:
            _rtl_para(cp1)
        crun = cp1.add_run()
        crun.text = card.card_title
        _style_run(crun, size=16, bold=True, color=theme_rgb["text"],
                   font=hdr["font"])

        # card body
        cbody_box = _add_textbox(slide, left + 0.22, top + 1.75,
                                 card_w - 0.44, card_h - 1.95)
        cbtf = cbody_box.text_frame
        cbtf.word_wrap = True
        cbp = cbtf.paragraphs[0]
        cbp.alignment = hdr["align"]
        if rtl:
            _rtl_para(cbp)
        cbdy = cbp.add_run()
        cbdy.text = card.card_body
        _style_run(cbdy, size=12, bold=False, color=theme_rgb["muted"],
                   font=hdr["font"])

    _slide_takeaway(slide, slide_data.key_takeaway, theme_rgb, rtl)
    _add_notes(slide, slide_data.speaker_notes)


def _add_process_slide(prs, layout, slide_data, theme_rgb, rtl):
    """Layout 4 — Numbered process / timeline for bullet slides."""
    slide = prs.slides.add_slide(layout)
    _set_bg(slide, _hex_to_rgb("#FFFFFF"))

    hdr = _slide_header(slide, slide_data.title, slide_data.subtitle, theme_rgb, rtl)

    bullets = slide_data.bullets or []
    n = min(len(bullets), 6)
    if n == 0:
        _slide_takeaway(slide, slide_data.key_takeaway, theme_rgb, rtl)
        _add_notes(slide, slide_data.speaker_notes)
        return

    top_start = 2.25
    bottom = 5.9
    available = bottom - top_start
    row_h = available / n

    # vertical timeline line
    line_left = 11.55 if rtl else 1.15
    _add_shape(
        slide, MSO_SHAPE.RECTANGLE,
        line_left, top_start + 0.05, 0.03, available - 0.1,
        fill_rgb=theme_rgb["accent_light"]
    )

    for idx, bullet in enumerate(bullets[:n]):
        row_top = top_start + idx * row_h

        # circle with number
        circle_left = line_left - 0.21
        _add_shape(
            slide, MSO_SHAPE.OVAL,
            circle_left, row_top, 0.45, 0.45,
            fill_rgb=theme_rgb["primary"]
        )
        num_box = _add_textbox(slide, circle_left, row_top + 0.03, 0.45, 0.4)
        np_ = num_box.text_frame.paragraphs[0]
        np_.alignment = PP_ALIGN.CENTER
        nrun = np_.add_run()
        nrun.text = str(idx + 1)
        _style_run(nrun, size=13, bold=True, color=_hex_to_rgb("#FFFFFF"),
                   font=hdr["font"])

        # text box
        if rtl:
            text_left = 0.7
            text_w = 10.6
        else:
            text_left = 1.75
            text_w = 10.85

        tbox = _add_textbox(slide, text_left, row_top - 0.02, text_w, row_h - 0.1)
        ttf = tbox.text_frame
        ttf.word_wrap = True
        tp = ttf.paragraphs[0]
        tp.alignment = hdr["align"]
        if rtl:
            _rtl_para(tp)
        trun = tp.add_run()
        trun.text = bullet
        _style_run(trun, size=15, bold=False, color=theme_rgb["text"],
                   font=hdr["font"])

    _slide_takeaway(slide, slide_data.key_takeaway, theme_rgb, rtl)
    _add_notes(slide, slide_data.speaker_notes)


def _add_highlight_slide(prs, layout, slide_data, theme_rgb, rtl):
    """Layout 5 — Large highlight statement for paragraph slides."""
    slide = prs.slides.add_slide(layout)
    _set_bg(slide, theme_rgb["soft_bg"])

    hdr = _slide_header(slide, slide_data.title, slide_data.subtitle, theme_rgb, rtl)

    # Big centered card
    card_left = 1.2
    card_top = 2.4
    card_w = 10.933
    card_h = 3.2

    _add_shape(
        slide, MSO_SHAPE.ROUNDED_RECTANGLE,
        card_left, card_top, card_w, card_h,
        fill_rgb=_hex_to_rgb("#FFFFFF"),
        line_rgb=theme_rgb["accent_light"],
        line_width=1.25
    )

    # accent left stripe
    stripe_left = card_left + card_w - 0.14 if rtl else card_left
    _add_shape(
        slide, MSO_SHAPE.ROUNDED_RECTANGLE,
        stripe_left, card_top + 0.3, 0.14, card_h - 0.6,
        fill_rgb=theme_rgb["primary"]
    )

    inner_left = card_left + 0.4
    inner_w = card_w - 0.8

    tbox = _add_textbox(slide, inner_left, card_top + 0.4, inner_w, card_h - 0.8)
    ttf = tbox.text_frame
    ttf.word_wrap = True
    tp = ttf.paragraphs[0]
    tp.alignment = PP_ALIGN.CENTER
    if rtl:
        _rtl_para(tp)
    trun = tp.add_run()
    trun.text = slide_data.paragraph_text or ""
    _style_run(trun, size=20, bold=False, color=theme_rgb["text"],
               font=hdr["font"])

    _slide_takeaway(slide, slide_data.key_takeaway, theme_rgb, rtl)
    _add_notes(slide, slide_data.speaker_notes)


def _add_example_slide(prs, layout, slide_data, theme_rgb, rtl):
    """Layout — Dark code-style example block."""
    slide = prs.slides.add_slide(layout)
    _set_bg(slide, _hex_to_rgb("#FFFFFF"))

    hdr = _slide_header(slide, slide_data.title, slide_data.subtitle, theme_rgb, rtl)

    # dark block
    block_left = 0.7
    block_top = 2.25
    block_w = 11.933
    block_h = 3.4

    _add_shape(
        slide, MSO_SHAPE.ROUNDED_RECTANGLE,
        block_left, block_top, block_w, block_h,
        fill_rgb=_hex_to_rgb("#0F172A")
    )

    # mac-style dots
    dot_colors = ["#EF4444", "#F59E0B", "#10B981"]
    for i, c in enumerate(dot_colors):
        dot_left = block_left + 0.35 + i * 0.28 if not rtl else block_left + block_w - 0.35 - i * 0.28
        _add_shape(
            slide, MSO_SHAPE.OVAL,
            dot_left, block_top + 0.28, 0.15, 0.15,
            fill_rgb=_hex_to_rgb(c)
        )

    # example title
    example_title = slide_data.example_title or ("Practical Example" if not rtl else "تطبيق عملي")
    title_left = block_left + 0.4
    etbox = _add_textbox(slide, title_left, block_top + 0.65,
                         block_w - 0.8, 0.5)
    etf = etbox.text_frame
    etf.word_wrap = True
    ep = etf.paragraphs[0]
    ep.alignment = hdr["align"]
    if rtl:
        _rtl_para(ep)
    erun = ep.add_run()
    erun.text = example_title
    _style_run(erun, size=15, bold=True,
               color=theme_rgb["primary"] if theme_rgb["key"] != "slate"
               else _hex_to_rgb("#93C5FD"),
               font=hdr["font"])

    # example body
    example_body = slide_data.example_body or (
        "Sample code or practical scenario example."
        if not rtl else
        "شرح سيناريو أو كود نموذجي."
    )
    bbox = _add_textbox(slide, title_left, block_top + 1.25,
                        block_w - 0.8, block_h - 1.5)
    btf = bbox.text_frame
    btf.word_wrap = True
    bp = btf.paragraphs[0]
    bp.alignment = hdr["align"]
    if rtl:
        _rtl_para(bp)
    brun = bp.add_run()
    brun.text = example_body
    mono = "Consolas" if not rtl else "Courier New"
    _style_run(brun, size=13, bold=False,
               color=_hex_to_rgb("#E2E8F0"),
               font=mono)

    _slide_takeaway(slide, slide_data.key_takeaway, theme_rgb, rtl)
    _add_notes(slide, slide_data.speaker_notes)


def _add_references_slide(prs, layout, slide_data, theme_rgb, rtl):
    """Layout — References slide with numbered sources."""
    slide = prs.slides.add_slide(layout)
    _set_bg(slide, _hex_to_rgb("#FFFFFF"))

    hdr = _slide_header(slide, slide_data.title, slide_data.subtitle, theme_rgb, rtl)

    refs = slide_data.references or []
    n = min(len(refs), 5)
    if n == 0:
        _slide_takeaway(slide, slide_data.key_takeaway, theme_rgb, rtl)
        _add_notes(slide, slide_data.speaker_notes)
        return

    top = 2.25
    row_h = 0.75

    for idx, ref in enumerate(refs[:n]):
        row_top = top + idx * row_h

        # subtle bg
        _add_shape(
            slide, MSO_SHAPE.ROUNDED_RECTANGLE,
            0.7, row_top, 11.933, row_h - 0.1,
            fill_rgb=theme_rgb["soft_bg"]
        )

        # number pill
        pill_left = 0.9 if not rtl else 11.5
        _add_shape(
            slide, MSO_SHAPE.ROUNDED_RECTANGLE,
            pill_left, row_top + 0.13, 0.55, 0.38,
            fill_rgb=theme_rgb["primary"]
        )
        pbox = _add_textbox(slide, pill_left, row_top + 0.15, 0.55, 0.36)
        pp = pbox.text_frame.paragraphs[0]
        pp.alignment = PP_ALIGN.CENTER
        prun = pp.add_run()
        prun.text = f"{idx + 1}"
        _style_run(prun, size=11, bold=True,
                   color=_hex_to_rgb("#FFFFFF"), font=hdr["font"])

        if rtl:
            text_left = 0.9
            text_w = 10.4
        else:
            text_left = 1.7
            text_w = 10.7

        # title + source
        tbox = _add_textbox(slide, text_left, row_top + 0.05, text_w, 0.65)
        ttf = tbox.text_frame
        ttf.word_wrap = True
        tp = ttf.paragraphs[0]
        tp.alignment = hdr["align"]
        if rtl:
            _rtl_para(tp)
        r1 = tp.add_run()
        r1.text = ref.title + "  "
        _style_run(r1, size=12, bold=True, color=theme_rgb["text"],
                   font=hdr["font"])
        r2 = tp.add_run()
        r2.text = f"— {ref.author_or_source}"
        _style_run(r2, size=11, italic=True, color=theme_rgb["primary"],
                   font=hdr["font"])

        dp = ttf.add_paragraph()
        dp.alignment = hdr["align"]
        if rtl:
            _rtl_para(dp)
        r3 = dp.add_run()
        r3.text = ref.description
        _style_run(r3, size=10, bold=False, color=theme_rgb["muted"],
                   font=hdr["font"])

    _slide_takeaway(slide, slide_data.key_takeaway, theme_rgb, rtl)
    _add_notes(slide, slide_data.speaker_notes)


def _add_notes(slide, notes_text):
    try:
        slide.notes_slide.notes_text_frame.text = notes_text or ""
    except Exception:
        pass


# ============================================================
# PPTX MAIN ENTRY
# ============================================================
def create_pptx(deck: SlideDeckSchema, filename: str, theme=None, rtl: bool = False):
    """
    Build a professionally designed PPTX using the given theme palette.
    Adds a hero title slide at the start and renders each content slide
    through a dedicated layout function based on slide.layout_type.
    """
    palette = dict(theme or THEMES["slate"])
    palette_rgb = {
        "key": palette.get("key", "slate"),
        "name": palette.get("name", "Slate"),
        "label": palette.get("label", "General"),
        "primary": _hex_to_rgb(palette["primary"]),
        "dark": _hex_to_rgb(palette["dark"]),
        "soft_bg": _hex_to_rgb(palette["soft_bg"]),
        "text": _hex_to_rgb(palette["text"]),
        "accent_light": _hex_to_rgb(palette["accent_light"]),
        "muted": _hex_to_rgb(palette["muted"]),
    }

    prs = PptxPresentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # 1) Hero title slide
    _add_title_slide(prs, blank_layout, deck.topic, palette_rgb, rtl)

    # 2) Content slides — dispatch by layout_type
    for slide_data in deck.slides:
        lt = slide_data.layout_type
        if lt == "cards" and slide_data.cards:
            _add_cards_slide(prs, blank_layout, slide_data, palette_rgb, rtl)
        elif lt == "bullets" and slide_data.bullets:
            _add_process_slide(prs, blank_layout, slide_data, palette_rgb, rtl)
        elif lt == "example":
            _add_example_slide(prs, blank_layout, slide_data, palette_rgb, rtl)
        elif lt == "references" and slide_data.references:
            _add_references_slide(prs, blank_layout, slide_data, palette_rgb, rtl)
        else:
            # paragraph + fallback
            _add_highlight_slide(prs, blank_layout, slide_data, palette_rgb, rtl)

    prs.save(filename)
    return filename


# ============================================================
# DOCX BUILDERS (unchanged behaviour)
# ============================================================
def create_assessment_docx(assessment: AssessmentSchema, filename: str, rtl: bool = False):
    """Generates a clean assessment document with optional Arabic/RTL alignment."""
    doc = Document()
    align_default = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT

    h0 = doc.add_heading(assessment.title, level=0)
    h0.alignment = align_default
    h0.runs[0].font.color.rgb = DocxRGBColor(24, 43, 73)

    p_inst = doc.add_paragraph(f"{'الإرشادات: ' if rtl else 'Instructions: '}{assessment.instructions}")
    p_inst.alignment = align_default
    doc.add_paragraph().paragraph_format.space_after = DocxPt(12)

    h1 = doc.add_heading("القسم الأول: أسئلة الاختيار من متعدد" if rtl else "Section 1: Multiple Choice Questions", level=1)
    h1.alignment = align_default

    for i, q in enumerate(assessment.mc_questions, start=1):
        qp = doc.add_paragraph()
        qp.alignment = align_default
        q_run = qp.add_run(f"س{i}. {q.question_text}" if rtl else f"Q{i}. {q.question_text}")
        q_run.bold = True
        qp.paragraph_format.space_before = DocxPt(8)
        for opt in q.options:
            op = doc.add_paragraph(opt)
            op.alignment = align_default
            op.paragraph_format.left_indent = DocxInches(0.0 if rtl else 0.4)
            op.paragraph_format.right_indent = DocxInches(0.4 if rtl else 0.0)
            op.paragraph_format.space_after = DocxPt(2)

    h2 = doc.add_heading("القسم الثاني: أسئلة الصواب والخطأ" if rtl else "Section 2: True or False Questions", level=1)
    h2.alignment = align_default

    for i, q in enumerate(assessment.tf_questions, start=11):
        qp = doc.add_paragraph()
        qp.alignment = align_default
        tf_prefix = "[صواب / خطأ] " if rtl else "[True / False] "
        q_run = qp.add_run(f"س{i}. {tf_prefix}{q.statement_text}" if rtl else f"Q{i}. {tf_prefix}{q.statement_text}")
        q_run.bold = True
        qp.paragraph_format.space_before = DocxPt(8)
        qp.paragraph_format.space_after = DocxPt(4)

    doc.add_page_break()
    hk = doc.add_heading("دليل الإجابات للمعلم" if rtl else "Teacher Answer Key", level=0)
    hk.alignment = align_default

    hk1 = doc.add_heading("إجابات قسم الاختيار من متعدد" if rtl else "Section 1: Multiple Choice Answers", level=2)
    hk1.alignment = align_default

    for i, q in enumerate(assessment.mc_questions, start=1):
        ak = doc.add_paragraph()
        ak.alignment = align_default
        ak.add_run(f"س{i} الإجابة الصحيحة: " if rtl else f"Q{i} Correct Answer: ").bold = True
        ak.add_run(f"{q.correct_answer}\n")
        exp = ak.add_run(f"الشرح: {q.explanation}" if rtl else f"Explanation: {q.explanation}")
        exp.font.italic = True
        ak.paragraph_format.space_after = DocxPt(6)

    hk2 = doc.add_heading("إجابات قسم الصواب والخطأ" if rtl else "Section 2: True / False Answers", level=2)
    hk2.alignment = align_default

    for i, q in enumerate(assessment.tf_questions, start=11):
        ak = doc.add_paragraph()
        ak.alignment = align_default
        ak.add_run(f"س{i} الإجابة الصحيحة: " if rtl else f"Q{i} Correct Answer: ").bold = True
        ak.add_run(f"{q.correct_answer}\n")
        exp = ak.add_run(f"الشرح: {q.explanation}" if rtl else f"Explanation: {q.explanation}")
        exp.font.italic = True
        ak.paragraph_format.space_after = DocxPt(6)

    doc.save(filename)
    return filename


def export_teaching_kit_docx(plan, teaching_kit, evaluation, sources, filename):
    doc = Document()
    doc.add_heading("SAGE Teaching Kit", level=0)
    doc.add_paragraph(f"Topic: {plan.topic}")
    doc.add_paragraph(f"Audience: {plan.audience}")
    doc.add_paragraph(f"Duration: {plan.duration}")

    doc.add_heading("Learning Objectives", level=1)
    for item in plan.objectives:
        doc.add_paragraph(item.objective, style="List Bullet")

    doc.add_heading("Key Concepts", level=1)
    for item in plan.key_concepts:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Teaching Kit", level=1)
    doc.add_heading("Summary", level=2)
    doc.add_paragraph(teaching_kit.summary)

    doc.add_heading("Key Terms", level=2)
    for item in teaching_kit.key_terms:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Explanation", level=2)
    for item in teaching_kit.explanation:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Examples", level=2)
    for item in teaching_kit.examples:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Teaching Tips", level=2)
    for item in teaching_kit.teaching_tips:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Common Misconceptions", level=2)
    for item in teaching_kit.misconceptions:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Evaluation", level=1)
    doc.add_paragraph(f"Final Score: {evaluation.score}/100")
    doc.add_paragraph(f"Approved: {evaluation.approved}")

    if evaluation.issues:
        doc.add_paragraph("Issues:")
        for issue in evaluation.issues:
            doc.add_paragraph(issue, style="List Bullet")

    if evaluation.improvements:
        doc.add_paragraph("Improvements:")
        for improvement in evaluation.improvements:
            doc.add_paragraph(improvement, style="List Bullet")

    doc.add_heading("References", level=1)
    seen = set()
    for source in sources:
        url = source.get("url", "")
        if url and url not in seen:
            seen.add(url)
            doc.add_paragraph(f"{source['title']} — {url}", style="List Bullet")

    doc.add_paragraph(
        "\nThe full 20-question assessment with the teacher answer key "
        "is provided in a separate Word document."
    )

    doc.save(filename)
    return filename


# ============================================================
# SESSION STATE INIT
# ============================================================
DEFAULTS = {
    "stage": "input",
    "runs_this_session": 0,
    "plan": None, "sources": None, "context": None,
    "teaching_kit": None, "presentation": None, "assessment": None,
    "evaluation": None, "revision_rounds": 0, "files": None,
    "rtl_mode": False,
    # Theme state
    "theme_mode": "auto",           # "auto" | "manual"
    "theme_key": "slate",           # emerald | tech_blue | indigo | amber | slate
    "theme_name": "General / Fallback",
    "accent_hex": "#475569",
    "accent_bg_hex": "#F8FAFC",
}
for k, v in DEFAULTS.items():
    if k not in st.session_state:
        st.session_state[k] = v


STAGES = [
    ("input", "Topic"),
    ("plan_review", "Plan"),
    ("teaching_kit_review", "Content"),
    ("presentation_review", "Presentation"),
    ("assessment_review", "Assessment"),
    ("evaluation_review", "Evaluation"),
    ("export_done", "Export"),
]


# ============================================================
# GLOBAL CSS
# ============================================================
CSS_TEMPLATE = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');

html, body, [class*="css"], .stMarkdown, .stTextInput, .stTextArea,
.stSelectbox, .stButton, .stRadio, .stProgress {
    font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
}

#MainMenu {visibility: hidden;}
footer {visibility: hidden;}
header[data-testid="stHeader"] {background: transparent;}

.stApp { background: #FAFBFC; }
.block-container {
    padding-top: 1.5rem !important;
    padding-bottom: 3rem !important;
    max-width: 1220px;
}

/* ============ TOP BAR ============ */
.sage-topbar {
    display: flex; align-items: center; justify-content: space-between;
    padding: 1rem 1.4rem;
    background: linear-gradient(135deg, #0F172A 0%, #1E293B 100%);
    border-radius: 14px; margin-bottom: 1rem; color: white;
    flex-wrap: wrap; gap: 0.8rem;
}
.sage-topbar-left { display: flex; flex-direction: column; gap: 2px; }
.sage-brand { font-size: 1.1rem; font-weight: 800; letter-spacing: 1px; color: white; }
.sage-brand-tagline { font-size: 0.8rem; color: #94A3B8; font-weight: 500; }
.sage-topbar-right { display: flex; gap: 0.5rem; align-items: center; flex-wrap: wrap; }
.sage-chip {
    display: inline-flex; align-items: center; gap: 0.4rem;
    background: rgba(255,255,255,0.08);
    border: 1px solid rgba(255,255,255,0.12);
    color: #E2E8F0; padding: 0.3rem 0.75rem;
    border-radius: 999px; font-size: 0.75rem; font-weight: 500;
    white-space: nowrap;
}
.sage-swatch {
    display: inline-block; width: 11px; height: 11px;
    border-radius: 3px; vertical-align: middle; margin-right: 2px;
}

/* ============ STEPPER ============ */
.sage-stepper {
    display: flex; align-items: center;
    padding: 1rem 1.25rem;
    background: white; border: 1px solid #E5E7EB;
    border-radius: 14px; margin-bottom: 1.5rem;
    overflow-x: auto;
}
.sage-step {
    display: flex; flex-direction: column; align-items: center;
    gap: 6px; min-width: 74px; flex-shrink: 0;
}
.sage-step-dot {
    width: 30px; height: 30px; border-radius: 50%;
    display: flex; align-items: center; justify-content: center;
    font-size: 12px; font-weight: 700;
    background: #F1F5F9; color: #94A3B8;
    border: 1.5px solid #E2E8F0;
    transition: all 0.2s ease;
}
.sage-step.active .sage-step-dot {
    background: __ACCENT__; color: white; border-color: __ACCENT__;
    box-shadow: 0 0 0 4px __ACCENT_BG__;
}
.sage-step.done .sage-step-dot {
    background: #10B981; color: white; border-color: #10B981;
}
.sage-step-label {
    font-size: 11.5px; color: #94A3B8; font-weight: 500;
    text-align: center; white-space: nowrap;
}
.sage-step.active .sage-step-label { color: #0F172A; font-weight: 700; }
.sage-step.done .sage-step-label { color: #059669; font-weight: 600; }
.sage-step-line {
    flex: 1; height: 2px; background: #E2E8F0;
    margin: 0 6px; margin-bottom: 24px;
    min-width: 16px; border-radius: 2px;
}
.sage-step-line.done { background: #10B981; }

/* ============ PAGE HEADER ============ */
.sage-page-header {
    margin-bottom: 1.2rem; padding-bottom: 1rem;
    border-bottom: 1px solid #E5E7EB;
}
.sage-page-title {
    font-size: 1.3rem; font-weight: 700; color: #0F172A; margin: 0 0 0.25rem 0;
}
.sage-page-subtitle {
    font-size: 0.9rem; color: #64748B; margin: 0; line-height: 1.5;
}

/* ============ SIDEBAR ============ */
section[data-testid="stSidebar"] {
    background: #0F172A; border-right: 1px solid #1E293B;
}
section[data-testid="stSidebar"] > div:first-child { padding-top: 1.5rem; }
section[data-testid="stSidebar"] * { color: #CBD5E1; }
section[data-testid="stSidebar"] hr { border-color: #1E293B; margin: 1rem 0; }
section[data-testid="stSidebar"] div.stButton > button {
    background: rgba(255,255,255,0.06); color: #E2E8F0 !important;
    border: 1px solid rgba(255,255,255,0.12);
    width: 100%; font-weight: 600; border-radius: 9px;
}
section[data-testid="stSidebar"] div.stButton > button:hover {
    background: rgba(255,255,255,0.12); border-color: rgba(255,255,255,0.22);
}
section[data-testid="stSidebar"] div[data-testid="stRadio"] label {
    color: #CBD5E1 !important; font-size: 0.85rem !important;
}
section[data-testid="stSidebar"] div[data-testid="stSelectbox"] label {
    color: #94A3B8 !important; font-size: 0.8rem !important;
}
section[data-testid="stSidebar"] div[data-testid="stSelectbox"] > div > div {
    background: #1E293B !important;
    border: 1px solid #334155 !important;
    color: #E2E8F0 !important;
    border-radius: 9px;
}
.sage-sidebar-brand {
    font-size: 1.35rem; font-weight: 800; color: #FFFFFF;
    letter-spacing: 1px; margin-bottom: 2px; display: block;
}
.sage-sidebar-tag {
    font-size: 0.78rem; color: #64748B;
    margin-bottom: 1rem; display: block; line-height: 1.4;
}
.sage-sidebar-label {
    font-size: 0.68rem; font-weight: 700; color: #64748B;
    letter-spacing: 1.2px; text-transform: uppercase;
    margin: 1.1rem 0 0.55rem 0;
}
.sage-side-step {
    display: flex; align-items: center; gap: 0.7rem;
    padding: 0.5rem 0.7rem; border-radius: 9px;
    margin-bottom: 3px; font-size: 0.86rem;
    color: #94A3B8; transition: all 0.15s ease;
}
.sage-side-step-icon {
    width: 22px; height: 22px; border-radius: 50%;
    display: flex; align-items: center; justify-content: center;
    font-size: 10px; font-weight: 700;
    background: #1E293B; color: #64748B; flex-shrink: 0;
}
.sage-side-step.active {
    background: rgba(255,255,255,0.07); color: #FFFFFF; font-weight: 700;
}
.sage-side-step.active .sage-side-step-icon {
    background: __ACCENT__; color: white;
    box-shadow: 0 0 0 3px rgba(255,255,255,0.08);
}
.sage-side-step.done .sage-side-step-icon { background: #10B981; color: white; }
.sage-side-step.done { color: #64748B; }
.sage-sidebar-footer {
    font-size: 0.7rem; color: #475569; text-align: center;
    margin-top: 1.5rem; letter-spacing: 0.4px;
}

/* ============ BUTTONS ============ */
div.stButton > button {
    border-radius: 9px; font-weight: 600; font-size: 0.9rem;
    padding: 0.5rem 1.15rem;
    border: 1px solid #E5E7EB; background: white; color: #0F172A;
    transition: all 0.15s ease;
}
div.stButton > button:hover { border-color: #CBD5E1; background: #F8FAFC; }
div.stButton > button[kind="primary"] {
    background: __ACCENT__; color: white; border: 1px solid __ACCENT__;
}
div.stButton > button[kind="primary"]:hover { filter: brightness(1.08); }

div[data-testid="stDownloadButton"] > button {
    border-radius: 9px; font-weight: 600;
    padding: 0.6rem 1rem; background: white;
    border: 1px solid #E5E7EB; color: #0F172A;
    width: 100%; transition: all 0.15s ease;
}
div[data-testid="stDownloadButton"] > button:hover {
    border-color: __ACCENT__; color: __ACCENT__;
}

/* ============ EXPANDERS ============ */
div[data-testid="stExpander"] {
    border-radius: 12px !important;
    border: 1px solid #E5E7EB !important;
    box-shadow: 0 1px 2px rgba(0,0,0,0.02);
    margin-bottom: 0.55rem; background: white; overflow: hidden;
}
div[data-testid="stExpander"] summary {
    padding: 0.75rem 1rem; font-weight: 600;
    font-size: 0.92rem; color: #0F172A;
}
div[data-testid="stExpander"] summary:hover { color: __ACCENT__; }

/* ============ CONTAINERS / METRICS ============ */
div[data-testid="stVerticalBlockBorderWrapper"] {
    border-radius: 14px !important; border-color: #E5E7EB !important;
}
div[data-testid="stMetric"] {
    background: white; border: 1px solid #E5E7EB;
    border-radius: 12px; padding: 0.9rem 1rem;
}
div[data-testid="stMetric"] label {
    font-size: 0.8rem; color: #64748B; font-weight: 500;
}
div[data-testid="stMetric"] div[data-testid="stMetricValue"] {
    font-size: 1.5rem; font-weight: 700; color: #0F172A;
}

/* ============ INPUTS ============ */
.stTextInput input, .stTextArea textarea {
    border-radius: 9px !important;
    border: 1px solid #E5E7EB !important;
    font-size: 0.92rem; background: white !important;
}
.stTextInput input:focus, .stTextArea textarea:focus {
    border-color: __ACCENT__ !important;
    box-shadow: 0 0 0 3px __ACCENT_BG__ !important;
}
.stTextArea label, .stTextInput label, .stSelectbox label, .stRadio label {
    font-size: 0.85rem !important;
    font-weight: 600 !important;
    color: #334155 !important;
}

/* ============ BADGES ============ */
.sage-badge {
    display: inline-flex; align-items: center; gap: 0.35rem;
    padding: 0.28rem 0.7rem; border-radius: 999px;
    font-size: 0.78rem; font-weight: 600;
}
.sage-badge.success { background:#ECFDF5; color:#047857; border:1px solid #A7F3D0; }
.sage-badge.warn    { background:#FFFBEB; color:#B45309; border:1px solid #FDE68A; }
.sage-badge.info    { background:#EFF6FF; color:#1D4ED8; border:1px solid #BFDBFE; }

/* ============ SOURCE ITEM ============ */
.sage-source-item {
    padding: 0.7rem 0.9rem; border-radius: 10px;
    background: #F8FAFC; border: 1px solid #E5E7EB;
    margin-bottom: 0.5rem; font-size: 0.85rem;
}
.sage-source-title { font-weight: 600; color: #0F172A; margin-bottom: 2px; }
.sage-source-meta { color: #64748B; font-size: 0.76rem; }

/* ============ SCORE PANEL ============ */
.sage-score-panel {
    display: flex; align-items: center; gap: 1.6rem;
    padding: 1.5rem 1.7rem; background: white;
    border: 1px solid #E5E7EB; border-radius: 14px;
    margin-bottom: 1.2rem; flex-wrap: wrap;
}
.sage-score-circle {
    width: 96px; height: 96px; border-radius: 50%;
    display: flex; flex-direction: column;
    align-items: center; justify-content: center;
    flex-shrink: 0; color: white;
    font-weight: 800; font-size: 1.7rem; line-height: 1;
}
.sage-score-circle small {
    font-size: 0.7rem; font-weight: 500; opacity: 0.85; margin-top: 3px;
}
.sage-score-info h3 {
    margin: 0 0 0.4rem 0; font-size: 1.1rem;
    color: #0F172A; font-weight: 700;
}
.sage-score-info p {
    margin: 0 0 0.5rem 0; color: #64748B; font-size: 0.88rem;
}

/* ============ SUCCESS HERO ============ */
.sage-success-hero {
    background: linear-gradient(135deg, #064E3B 0%, #10B981 100%);
    border-radius: 16px; padding: 2rem 2.2rem;
    color: white; margin-bottom: 1.5rem;
    box-shadow: 0 4px 16px rgba(16,185,129,0.18);
}
.sage-success-hero h2 {
    margin: 0 0 0.5rem 0; font-size: 1.55rem;
    font-weight: 800; color: white; letter-spacing: -0.3px;
}
.sage-success-hero p {
    margin: 0; color: rgba(255,255,255,0.9);
    font-size: 0.95rem; line-height: 1.55;
}

/* ============ LISTS ============ */
.sage-list { list-style: none; padding: 0; margin: 0.4rem 0 1rem 0; }
.sage-list li {
    padding: 0.6rem 0.85rem; background: #F8FAFC;
    border: 1px solid #E5E7EB; border-radius: 9px;
    margin-bottom: 0.4rem; font-size: 0.88rem;
    color: #334155; line-height: 1.5;
}

/* ============ HERO (INPUT) ============ */
.sage-hero {
    max-width: 720px; margin: 0 auto;
    text-align: center; padding: 0 1rem;
}
.sage-hero-logo {
    width: 64px; height: 64px; border-radius: 18px;
    background: __ACCENT_BG__; color: __ACCENT__;
    display: flex; align-items: center; justify-content: center;
    font-size: 2rem; font-weight: 700;
    margin: 0 auto 1.4rem auto;
    border: 1px solid __ACCENT__22;
}
.sage-hero-title {
    font-size: 2.3rem; font-weight: 800; color: #0F172A;
    letter-spacing: -0.6px; margin: 0 0 0.5rem 0;
    line-height: 1.15;
}
.sage-hero-tagline {
    font-size: 0.9rem; font-weight: 600;
    color: __ACCENT__; letter-spacing: 1.2px;
    text-transform: uppercase; margin-bottom: 1.1rem;
}
.sage-hero-sub {
    font-size: 1.02rem; color: #64748B;
    line-height: 1.65; margin: 0 0 2rem 0;
}
.sage-feature-row {
    display: flex; justify-content: center; gap: 2rem;
    margin-top: 2.4rem; flex-wrap: wrap;
}
.sage-feature {
    display: flex; flex-direction: column;
    align-items: center; gap: 0.55rem;
    font-size: 0.8rem; color: #94A3B8; font-weight: 500;
}
.sage-feature-icon {
    width: 42px; height: 42px; border-radius: 11px;
    background: #F1F5F9; color: #64748B;
    display: flex; align-items: center; justify-content: center;
    font-size: 0.95rem; font-weight: 700;
}

/* ============ THEME PICKER (INPUT) ============ */
.sage-theme-strip {
    display: flex; justify-content: center;
    gap: 0.7rem; margin: 1.5rem 0 0 0;
    flex-wrap: wrap;
}
.sage-theme-pill {
    display: inline-flex; align-items: center; gap: 0.45rem;
    padding: 0.45rem 0.9rem;
    background: white; border: 1px solid #E5E7EB;
    border-radius: 999px; font-size: 0.82rem;
    color: #334155; font-weight: 500;
}
.sage-theme-pill.active {
    border-color: __ACCENT__;
    color: __ACCENT__; font-weight: 700;
    box-shadow: 0 0 0 3px __ACCENT_BG__;
}
.sage-theme-dot {
    width: 12px; height: 12px; border-radius: 50%;
}

/* ============ DIVIDER ============ */
hr { border-color: #E5E7EB; margin: 1.5rem 0; }

/* ============ RESPONSIVE ============ */
@media (max-width: 768px) {
    .sage-hero-title { font-size: 1.8rem; }
    .sage-feature-row { gap: 1rem; }
    .sage-score-panel { gap: 1rem; padding: 1.1rem; }
    .sage-score-circle { width: 78px; height: 78px; font-size: 1.4rem; }
}
</style>
"""


def inject_global_css(accent: str, accent_bg: str):
    css = CSS_TEMPLATE.replace("__ACCENT__", accent).replace("__ACCENT_BG__", accent_bg)
    st.markdown(css, unsafe_allow_html=True)


# ============================================================
# SIDEBAR
# ============================================================
def render_sidebar():
    with st.sidebar:
        st.markdown(
            "<span class='sage-sidebar-brand'>✦ SAGE</span>"
            "<span class='sage-sidebar-tag'>Agentic AI Teaching Kit Generator</span>",
            unsafe_allow_html=True
        )

        current_idx = next(
            (i for i, (k, _) in enumerate(STAGES) if k == st.session_state.stage), 0
        )
        total = len(STAGES)

        st.markdown(
            f"<div class='sage-sidebar-label'>Progress · {current_idx + 1} / {total}</div>",
            unsafe_allow_html=True
        )
        st.progress(current_idx / (total - 1) if total > 1 else 0)

        st.markdown("<div class='sage-sidebar-label'>Workflow</div>", unsafe_allow_html=True)
        for i, (key, label) in enumerate(STAGES):
            if key == st.session_state.stage:
                cls, mark = "active", str(i + 1)
            elif i < current_idx:
                cls, mark = "done", "✓"
            else:
                cls, mark = "", str(i + 1)

            st.markdown(
                f"<div class='sage-side-step {cls}'>"
                f"<div class='sage-side-step-icon'>{mark}</div>"
                f"<span>{label}</span>"
                f"</div>",
                unsafe_allow_html=True
            )

        # ------- Theme settings -------
        st.markdown("<div class='sage-sidebar-label'>PowerPoint Theme</div>",
                    unsafe_allow_html=True)

        st.radio(
            "Theme mode",
            options=["auto", "manual"],
            key="theme_mode",
            format_func=lambda x: "Auto-detect from topic" if x == "auto" else "Manual selection",
            label_visibility="collapsed"
        )

        theme_keys = ["emerald", "tech_blue", "indigo", "amber", "slate"]
        if st.session_state.theme_mode == "manual":
            st.selectbox(
                "Theme",
                options=theme_keys,
                key="theme_key",
                format_func=lambda k: THEMES[k]["label"],
                label_visibility="collapsed"
            )
        else:
            # Auto mode — detect from plan topic if available
            if st.session_state.plan is not None:
                st.session_state.theme_key = detect_theme(st.session_state.plan.topic)

        # Sync accent colors + name based on current theme_key
        _pal = get_theme_palette(st.session_state.theme_key)
        st.session_state.theme_name = THEMES[st.session_state.theme_key]["label"]
        st.session_state.accent_hex = _pal["primary"]
        st.session_state.accent_bg_hex = _pal["soft_bg"]

        # Preview
        st.markdown(
            f"""
            <div style='margin-top:0.6rem;'>
                <div style='display:flex;gap:8px;align-items:center;'>
                    <div style='width:26px;height:26px;border-radius:7px;background:{_pal["primary"]};'></div>
                    <div style='width:26px;height:26px;border-radius:7px;background:{_pal["soft_bg"]};border:1px solid #334155;'></div>
                    <div style='width:26px;height:26px;border-radius:7px;background:{_pal["dark"]};'></div>
                </div>
                <div style='font-size:0.72rem;color:#94A3B8;margin-top:8px;line-height:1.4;'>
                    <b style='color:#E2E8F0;'>{THEMES[st.session_state.theme_key]["name"]}</b><br>
                    {st.session_state.theme_name}
                </div>
            </div>
            """,
            unsafe_allow_html=True
        )

        st.markdown("---")

        if st.button("Start new session", use_container_width=True):
            for k, v in DEFAULTS.items():
                if k != "runs_this_session":
                    st.session_state[k] = v
            st.rerun()

        st.markdown(
            "<div class='sage-sidebar-footer'>Powered by SAGE Engine</div>",
            unsafe_allow_html=True
        )


def render_topbar():
    theme = THEMES.get(st.session_state.theme_key, THEMES["slate"])
    lang_label = "AR · RTL" if st.session_state.rtl_mode else "EN"
    st.markdown(
        f"""
        <div class='sage-topbar'>
            <div class='sage-topbar-left'>
                <div class='sage-brand'>✦ SAGE</div>
                <div class='sage-brand-tagline'>Agentic AI Teaching Kit Generator</div>
            </div>
            <div class='sage-topbar-right'>
                <span class='sage-chip'>
                    <span class='sage-swatch' style='background:{theme["primary"]};'></span>
                    {theme["name"]}
                </span>
                <span class='sage-chip'>{lang_label}</span>
            </div>
        </div>
        """,
        unsafe_allow_html=True
    )


def render_stepper():
    current_idx = next(
        (i for i, (k, _) in enumerate(STAGES) if k == st.session_state.stage), 0
    )
    html = "<div class='sage-stepper'>"
    for i, (_key, label) in enumerate(STAGES):
        if i < current_idx:
            cls, mark = "done", "✓"
        elif i == current_idx:
            cls, mark = "active", str(i + 1)
        else:
            cls, mark = "upcoming", str(i + 1)
        html += (
            f"<div class='sage-step {cls}'>"
            f"<div class='sage-step-dot'>{mark}</div>"
            f"<div class='sage-step-label'>{label}</div>"
            f"</div>"
        )
        if i < len(STAGES) - 1:
            line_cls = "done" if i < current_idx else ""
            html += f"<div class='sage-step-line {line_cls}'></div>"
    html += "</div>"
    st.markdown(html, unsafe_allow_html=True)


def render_page_header(title: str, subtitle: str):
    st.markdown(
        f"<div class='sage-page-header'>"
        f"<div class='sage-page-title'>{title}</div>"
        f"<p class='sage-page-subtitle'>{subtitle}</p>"
        f"</div>",
        unsafe_allow_html=True
    )


# ============================================================
# APPLY SIDEBAR FIRST (updates theme state) THEN CSS
# ============================================================
render_sidebar()
inject_global_css(st.session_state.accent_hex, st.session_state.accent_bg_hex)
render_topbar()
render_stepper()


# ============================================================
# STAGE: INPUT
# ============================================================
if st.session_state.stage == "input":

    if st.session_state.runs_this_session >= MAX_SESSION_RUNS:
        st.warning(
            f"You have reached the maximum number of runs for this session "
            f"({MAX_SESSION_RUNS}). Reload the page to continue."
        )
    else:
        # Vertical spacer to push hero into the visual middle of the viewport
        st.markdown("<div style='height: 10vh;'></div>", unsafe_allow_html=True)

        _l, center, _r = st.columns([1, 2, 1])
        with center:
            st.markdown(
                f"""
                <div class='sage-hero'>
                    <div class='sage-hero-logo'>✦</div>
                    <div class='sage-hero-tagline'>Agentic AI Teaching Kit Generator</div>
                    <h1 class='sage-hero-title'>What would you like to teach today?</h1>
                    <p class='sage-hero-sub'>
                        Generate a structured lesson plan, teaching content, presentation,
                        assessment, evaluation, and export package — all from a single topic.
                    </p>
                </div>
                """,
                unsafe_allow_html=True
            )

            topic = st.text_input(
                "topic_input",
                placeholder="e.g. Fundamentals of Cybersecurity",
                label_visibility="collapsed"
            )

            generate = st.button(
                "Generate Teaching Kit",
                type="primary",
                disabled=not topic.strip(),
                use_container_width=True
            )

            # Workflow preview chips
            st.markdown(
                """
                <div class='sage-feature-row'>
                    <div class='sage-feature'><div class='sage-feature-icon'>1</div><span>Plan</span></div>
                    <div class='sage-feature'><div class='sage-feature-icon'>2</div><span>Research</span></div>
                    <div class='sage-feature'><div class='sage-feature-icon'>3</div><span>Design</span></div>
                    <div class='sage-feature'><div class='sage-feature-icon'>4</div><span>Evaluate</span></div>
                    <div class='sage-feature'><div class='sage-feature-icon'>5</div><span>Export</span></div>
                </div>
                """,
                unsafe_allow_html=True
            )

        if generate:
            if not daily_usage_ok_and_increment():
                st.error("Daily usage limit reached for this site. Please try again later.")
            else:
                with st.spinner("Generating the lesson plan..."):
                    try:
                        clean_topic = topic.strip()
                        rtl_mode = is_arabic(clean_topic)

                        # Theme: auto-detect (or keep manual selection)
                        if st.session_state.theme_mode == "auto":
                            st.session_state.theme_key = detect_theme(clean_topic)
                        _pal = get_theme_palette(st.session_state.theme_key)
                        st.session_state.theme_name = THEMES[st.session_state.theme_key]["label"]
                        st.session_state.accent_hex = _pal["primary"]
                        st.session_state.accent_bg_hex = _pal["soft_bg"]

                        st.session_state.rtl_mode = rtl_mode
                        st.session_state.plan = planner_agent(clean_topic)
                        st.session_state.runs_this_session += 1
                        st.session_state.stage = "plan_review"
                        st.rerun()
                    except Exception as e:
                        st.error(f"An error occurred while generating the plan: {e}")


# ============================================================
# STAGE: PLAN REVIEW
# ============================================================
elif st.session_state.stage == "plan_review":
    plan = st.session_state.plan

    render_page_header(
        "Review the lesson plan",
        "Review the plan generated by the agent. Edit any field before continuing to research."
    )

    with st.container(border=True):
        st.markdown("##### General information")
        c1, c2, c3 = st.columns(3)
        edit_topic = c1.text_input("Topic", value=plan.topic)
        edit_audience = c2.text_input("Audience", value=plan.audience)
        edit_duration = c3.text_input("Duration", value=plan.duration)

    col_left, col_right = st.columns(2)
    with col_left:
        with st.container(border=True):
            st.markdown("##### Learning objectives")
            st.caption("One objective per line.")
            edit_objectives = st.text_area(
                "objectives",
                value=list_to_text([o.objective for o in plan.objectives]),
                height=160, label_visibility="collapsed"
            )
        with st.container(border=True):
            st.markdown("##### Lesson flow")
            st.caption("One step per line.")
            edit_flow = st.text_area(
                "flow",
                value=list_to_text(plan.lesson_flow),
                height=160, label_visibility="collapsed"
            )
    with col_right:
        with st.container(border=True):
            st.markdown("##### Key concepts")
            st.caption("One concept per line.")
            edit_concepts = st.text_area(
                "concepts",
                value=list_to_text(plan.key_concepts),
                height=358, label_visibility="collapsed"
            )

    st.write("")
    col1, col2, _spacer = st.columns([1, 1, 3])
    with col1:
        if st.button("Confirm & research", type="primary", use_container_width=True):
            plan.topic = edit_topic
            plan.audience = edit_audience
            plan.duration = edit_duration
            plan.objectives = [Objective(objective=o) for o in text_to_list(edit_objectives)]
            plan.key_concepts = text_to_list(edit_concepts)
            plan.lesson_flow = text_to_list(edit_flow)
            st.session_state.plan = plan

            # Auto-detect theme based on (possibly edited) topic
            if st.session_state.theme_mode == "auto":
                st.session_state.theme_key = detect_theme(plan.topic)
                _pal = get_theme_palette(st.session_state.theme_key)
                st.session_state.theme_name = THEMES[st.session_state.theme_key]["label"]
                st.session_state.accent_hex = _pal["primary"]
                st.session_state.accent_bg_hex = _pal["soft_bg"]

            with st.spinner("Retrieving trusted sources (Wikipedia + OpenAlex) and building content..."):
                try:
                    sources = research_agent(plan)
                    context = build_context(sources)
                    teaching_kit = teaching_kit_agent(plan, context)
                    st.session_state.sources = sources
                    st.session_state.context = context
                    st.session_state.teaching_kit = teaching_kit
                    st.session_state.stage = "teaching_kit_review"
                    st.rerun()
                except Exception as e:
                    st.error(f"An error occurred during research/content generation: {e}")

    with col2:
        if st.button("Back to start", use_container_width=True):
            st.session_state.stage = "input"
            st.rerun()


# ============================================================
# STAGE: TEACHING KIT REVIEW
# ============================================================
elif st.session_state.stage == "teaching_kit_review":
    tk = st.session_state.teaching_kit
    sources = st.session_state.sources

    render_page_header(
        "Review teaching content",
        "Review the content generated from the retrieved sources. Edit any section before continuing."
    )

    with st.expander(f"Retrieved sources ({len(sources)} excerpts)", expanded=False):
        seen = set()
        for s in sources:
            if s["url"] and s["url"] not in seen:
                seen.add(s["url"])
                st.markdown(
                    f"<div class='sage-source-item'>"
                    f"<div class='sage-source-title'>{s['title']}</div>"
                    f"<div class='sage-source-meta'>{s['type']} · {s['url']}</div>"
                    f"</div>",
                    unsafe_allow_html=True
                )

    with st.container(border=True):
        st.markdown("##### Summary")
        edit_summary = st.text_area("summary", value=tk.summary, height=100, label_visibility="collapsed")

    c1, c2 = st.columns(2)
    with c1:
        with st.container(border=True):
            st.markdown("##### Key terms")
            st.caption("One term per line.")
            edit_terms = st.text_area("terms", value=list_to_text(tk.key_terms),
                                      height=140, label_visibility="collapsed")
        with st.container(border=True):
            st.markdown("##### Examples")
            st.caption("One example per line.")
            edit_examples = st.text_area("examples", value=list_to_text(tk.examples),
                                         height=140, label_visibility="collapsed")
        with st.container(border=True):
            st.markdown("##### Common misconceptions")
            st.caption("One misconception per line.")
            edit_misconceptions = st.text_area("misc", value=list_to_text(tk.misconceptions),
                                               height=140, label_visibility="collapsed")
    with c2:
        with st.container(border=True):
            st.markdown("##### Explanation")
            st.caption("One point per line.")
            edit_explanation = st.text_area("expl", value=list_to_text(tk.explanation),
                                            height=220, label_visibility="collapsed")
        with st.container(border=True):
            st.markdown("##### Teaching tips")
            st.caption("One tip per line.")
            edit_tips = st.text_area("tips", value=list_to_text(tk.teaching_tips),
                                     height=220, label_visibility="collapsed")

    st.write("")
    col1, col2, _spacer = st.columns([1, 1, 3])
    with col1:
        if st.button("Confirm & build presentation", type="primary", use_container_width=True):
            tk.summary = edit_summary
            tk.key_terms = text_to_list(edit_terms)
            tk.explanation = text_to_list(edit_explanation)
            tk.examples = text_to_list(edit_examples)
            tk.teaching_tips = text_to_list(edit_tips)
            tk.misconceptions = text_to_list(edit_misconceptions)
            st.session_state.teaching_kit = tk

            with st.spinner("Generating presentation slides (may take a minute)..."):
                try:
                    presentation = presentation_agent(
                        st.session_state.plan, tk, st.session_state.context,
                        rtl_mode=st.session_state.rtl_mode
                    )
                    st.session_state.presentation = presentation
                    st.session_state.stage = "presentation_review"
                    st.rerun()
                except Exception as e:
                    st.error(f"An error occurred while generating the presentation: {e}")

    with col2:
        if st.button("Back to plan review", use_container_width=True):
            st.session_state.stage = "plan_review"
            st.rerun()


# ============================================================
# STAGE: PRESENTATION REVIEW
# ============================================================
elif st.session_state.stage == "presentation_review":
    deck = st.session_state.presentation

    render_page_header(
        "Review presentation slides",
        "Review the 10 generated slides. A references slide is appended automatically from real sources at export time."
    )

    layout_options = ["cards", "bullets", "paragraph", "example"]

    for i, slide in enumerate(deck.slides):
        with st.expander(f"Slide {i + 1} · {slide.title}", expanded=False):
            c1, c2 = st.columns([2, 1])
            new_title = c1.text_input("Title", value=slide.title, key=f"s{i}_title")
            new_subtitle = c2.text_input("Subtitle", value=slide.subtitle, key=f"s{i}_subtitle")

            new_layout = st.selectbox(
                "Layout type", layout_options,
                index=layout_options.index(slide.layout_type) if slide.layout_type in layout_options else 0,
                key=f"s{i}_layout"
            )

            st.caption("The following fields are only used when they match the selected layout type above.")

            existing_cards = slide.cards or []
            card_titles, card_bodies = [], []
            for c_idx in range(3):
                cc1, cc2 = st.columns([1, 2])
                default_title = existing_cards[c_idx].card_title if c_idx < len(existing_cards) else ""
                default_body = existing_cards[c_idx].card_body if c_idx < len(existing_cards) else ""
                card_titles.append(cc1.text_input(f"Card {c_idx + 1} title",
                                                  value=default_title,
                                                  key=f"s{i}_card{c_idx}_title"))
                card_bodies.append(cc2.text_input(f"Card {c_idx + 1} body",
                                                  value=default_body,
                                                  key=f"s{i}_card{c_idx}_body"))

            new_bullets = st.text_area("Bullets (one per line)",
                                       value=list_to_text(slide.bullets or []),
                                       height=100, key=f"s{i}_bullets")
            new_paragraph = st.text_area("Paragraph text",
                                         value=slide.paragraph_text or "",
                                         height=100, key=f"s{i}_paragraph")
            new_example_title = st.text_input("Example title",
                                              value=slide.example_title or "",
                                              key=f"s{i}_ex_title")
            new_example_body = st.text_area("Example / code content",
                                            value=slide.example_body or "",
                                            height=100, key=f"s{i}_ex_body")

            new_takeaway = st.text_input("Key takeaway", value=slide.key_takeaway, key=f"s{i}_takeaway")
            new_notes = st.text_area("Speaker notes", value=slide.speaker_notes,
                                     height=100, key=f"s{i}_notes")

    st.write("")
    col1, col2, _spacer = st.columns([1, 1, 3])
    with col1:
        if st.button("Confirm & build assessment", type="primary", use_container_width=True):
            new_slides = []
            for i in range(len(deck.slides)):
                layout = st.session_state[f"s{i}_layout"]
                cards = None
                bullets = None
                paragraph_text = None
                example_title = None
                example_body = None

                if layout == "cards":
                    cards = [
                        SlideCard(
                            card_title=st.session_state[f"s{i}_card{c}_title"],
                            card_body=st.session_state[f"s{i}_card{c}_body"]
                        )
                        for c in range(3)
                        if st.session_state[f"s{i}_card{c}_title"].strip()
                    ]
                elif layout == "bullets":
                    bullets = text_to_list(st.session_state[f"s{i}_bullets"])
                elif layout == "paragraph":
                    paragraph_text = st.session_state[f"s{i}_paragraph"]
                elif layout == "example":
                    example_title = st.session_state[f"s{i}_ex_title"]
                    example_body = st.session_state[f"s{i}_ex_body"]

                new_slides.append(Slide(
                    title=st.session_state[f"s{i}_title"],
                    subtitle=st.session_state[f"s{i}_subtitle"],
                    layout_type=layout,
                    cards=cards, bullets=bullets, paragraph_text=paragraph_text,
                    example_title=example_title, example_body=example_body,
                    references=None,
                    key_takeaway=st.session_state[f"s{i}_takeaway"],
                    speaker_notes=st.session_state[f"s{i}_notes"],
                ))

            deck.slides = new_slides
            st.session_state.presentation = deck

            with st.spinner("Generating assessment (20 questions)..."):
                try:
                    assessment = assessment_agent(
                        st.session_state.plan, st.session_state.teaching_kit,
                        st.session_state.context,
                        rtl_mode=st.session_state.rtl_mode
                    )
                    st.session_state.assessment = assessment
                    st.session_state.stage = "assessment_review"
                    st.rerun()
                except Exception as e:
                    st.error(f"An error occurred while generating the assessment: {e}")

    with col2:
        if st.button("Back to content review", use_container_width=True):
            st.session_state.stage = "teaching_kit_review"
            st.rerun()


# ============================================================
# STAGE: ASSESSMENT REVIEW
# ============================================================
elif st.session_state.stage == "assessment_review":
    assessment = st.session_state.assessment

    render_page_header(
        "Review assessment",
        "Review the 20 questions (10 multiple-choice + 10 true/false). You can edit any question or option."
    )

    with st.container(border=True):
        edit_title = st.text_input("Assessment title", value=assessment.title)
        edit_instructions = st.text_area("Instructions", value=assessment.instructions, height=80)

    st.markdown("#### Section 1 · Multiple choice questions")
    for i, q in enumerate(assessment.mc_questions):
        with st.expander(f"Question {i + 1} · {q.question_text[:60]}", expanded=False):
            st.text_input("Question text", value=q.question_text, key=f"mc{i}_q")
            opts = []
            for o_idx in range(4):
                default_opt = q.options[o_idx] if o_idx < len(q.options) else ""
                opts.append(st.text_input(f"Option {o_idx + 1}",
                                          value=default_opt,
                                          key=f"mc{i}_opt{o_idx}"))
            default_correct_idx = q.options.index(q.correct_answer) if q.correct_answer in q.options else 0
            st.selectbox(
                "Correct answer",
                options=list(range(4)),
                index=default_correct_idx,
                format_func=lambda idx, _i=i: st.session_state.get(f"mc{_i}_opt{idx}", f"Option {idx+1}"),
                key=f"mc{i}_correct_idx"
            )
            st.text_area("Explanation", value=q.explanation, height=70, key=f"mc{i}_exp")

    st.markdown("#### Section 2 · True / False questions")
    for i, q in enumerate(assessment.tf_questions):
        with st.expander(f"Question {i + 11} · {q.statement_text[:60]}", expanded=False):
            st.text_input("Statement", value=q.statement_text, key=f"tf{i}_statement")
            st.radio(
                "Correct answer",
                ["True", "False"],
                index=0 if q.correct_answer == "True" else 1,
                key=f"tf{i}_correct",
                horizontal=True
            )
            st.text_area("Explanation", value=q.explanation, height=70, key=f"tf{i}_exp")

    st.write("")
    col1, col2, _spacer = st.columns([1, 1, 3])
    with col1:
        if st.button("Confirm & evaluate", type="primary", use_container_width=True):
            new_mc = []
            for i in range(len(assessment.mc_questions)):
                opts = [st.session_state[f"mc{i}_opt{o}"] for o in range(4)]
                correct_idx = st.session_state[f"mc{i}_correct_idx"]
                new_mc.append(MultipleChoiceQuestion(
                    question_text=st.session_state[f"mc{i}_q"],
                    options=opts,
                    correct_answer=opts[correct_idx],
                    explanation=st.session_state[f"mc{i}_exp"],
                ))

            new_tf = []
            for i in range(len(assessment.tf_questions)):
                new_tf.append(TrueFalseQuestion(
                    statement_text=st.session_state[f"tf{i}_statement"],
                    correct_answer=st.session_state[f"tf{i}_correct"],
                    explanation=st.session_state[f"tf{i}_exp"],
                ))

            assessment.title = edit_title
            assessment.instructions = edit_instructions
            assessment.mc_questions = new_mc
            assessment.tf_questions = new_tf
            st.session_state.assessment = assessment

            with st.spinner("Evaluating the full teaching kit..."):
                try:
                    evaluation = evaluator_agent(
                        st.session_state.plan, st.session_state.teaching_kit,
                        st.session_state.presentation, assessment,
                        st.session_state.context
                    )
                    st.session_state.evaluation = evaluation
                    st.session_state.stage = "evaluation_review"
                    st.rerun()
                except Exception as e:
                    st.error(f"An error occurred during evaluation: {e}")

    with col2:
        if st.button("Back to presentation review", use_container_width=True):
            st.session_state.stage = "presentation_review"
            st.rerun()


# ============================================================
# STAGE: EVALUATION REVIEW
# ============================================================
elif st.session_state.stage == "evaluation_review":
    ev = st.session_state.evaluation

    render_page_header(
        "Final evaluation result",
        "The agent evaluated the full kit. Export directly, or run auto-revision if it was not approved."
    )

    if ev.score >= 85:
        score_color = "#10B981"
    elif ev.score >= 70:
        score_color = "#F59E0B"
    else:
        score_color = "#EF4444"

    status_label = "Approved" if ev.approved else "Needs revision"
    status_cls = "success" if ev.approved else "warn"

    st.markdown(
        f"""
        <div class='sage-score-panel'>
            <div class='sage-score-circle' style='background:{score_color};'>
                {ev.score}<small>/ 100</small>
            </div>
            <div class='sage-score-info'>
                <h3>Overall quality score</h3>
                <p>Evaluated on factual accuracy, source grounding, objective alignment, presentation quality, and assessment quality.</p>
                <span class='sage-badge {status_cls}'>{status_label}</span>
            </div>
        </div>
        """,
        unsafe_allow_html=True
    )

    # Selected theme preview (informational)
    _t = THEMES.get(st.session_state.theme_key, THEMES["slate"])
    st.markdown(
        f"""
        <div style='display:flex;align-items:center;gap:0.7rem;padding:0.7rem 0.9rem;
                    background:white;border:1px solid #E5E7EB;border-radius:12px;margin-bottom:1.2rem;'>
            <span style='font-size:0.8rem;color:#64748B;font-weight:600;letter-spacing:0.4px;
                         text-transform:uppercase;'>PowerPoint Theme</span>
            <span class='sage-swatch' style='background:{_t["primary"]};width:14px;height:14px;'></span>
            <span class='sage-swatch' style='background:{_t["soft_bg"]};width:14px;height:14px;border:1px solid #E5E7EB;'></span>
            <span style='font-weight:600;color:#0F172A;font-size:0.9rem;'>{_t["name"]}</span>
            <span style='color:#64748B;font-size:0.85rem;'>· {_t["label"]}</span>
        </div>
        """,
        unsafe_allow_html=True
    )

    c1, c2 = st.columns(2)
    with c1:
        with st.container(border=True):
            st.markdown("##### Issues detected")
            if ev.issues:
                html_items = "".join(f"<li>{issue}</li>" for issue in ev.issues)
                st.markdown(f"<ul class='sage-list'>{html_items}</ul>", unsafe_allow_html=True)
            else:
                st.markdown("<p style='color:#64748B;font-size:0.88rem;'>No issues detected.</p>",
                            unsafe_allow_html=True)
    with c2:
        with st.container(border=True):
            st.markdown("##### Suggested improvements")
            if ev.improvements:
                html_items = "".join(f"<li>{imp}</li>" for imp in ev.improvements)
                st.markdown(f"<ul class='sage-list'>{html_items}</ul>", unsafe_allow_html=True)
            else:
                st.markdown("<p style='color:#64748B;font-size:0.88rem;'>No additional improvements suggested.</p>",
                            unsafe_allow_html=True)

    st.write("")
    col1, col2, col3 = st.columns(3)

    with col1:
        if st.button("Export files now", type="primary", use_container_width=True):
            with st.spinner("Building final files..."):
                try:
                    plan = st.session_state.plan
                    presentation = st.session_state.presentation
                    sources = st.session_state.sources

                    rtl_mode = st.session_state.rtl_mode
                    theme_palette = get_theme_palette(st.session_state.theme_key)

                    presentation.slides = [
                        s for s in presentation.slides if s.layout_type != "references"
                    ]
                    presentation.slides.append(build_references_slide(sources, rtl_mode=rtl_mode))

                    tmp_dir = tempfile.mkdtemp()
                    topic_title = sanitize_filename(plan.topic)

                    pptx_path = os.path.join(tmp_dir, f"{topic_title}_Presentation.pptx")
                    kit_docx_path = os.path.join(tmp_dir, f"{topic_title}_Teaching_Kit.docx")
                    assessment_docx_path = os.path.join(tmp_dir, f"{topic_title}_Assessment.docx")
                    zip_path = os.path.join(tmp_dir, f"{topic_title}_SAGE_Package.zip")

                    create_pptx(
                        presentation, pptx_path,
                        theme=theme_palette, rtl=rtl_mode
                    )
                    export_teaching_kit_docx(plan, st.session_state.teaching_kit, ev, sources, kit_docx_path)
                    create_assessment_docx(st.session_state.assessment, assessment_docx_path, rtl=rtl_mode)

                    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
                        archive.write(pptx_path, arcname=os.path.basename(pptx_path))
                        archive.write(kit_docx_path, arcname=os.path.basename(kit_docx_path))
                        archive.write(assessment_docx_path, arcname=os.path.basename(assessment_docx_path))

                    st.session_state.files = {
                        "pptx": pptx_path, "kit_docx": kit_docx_path,
                        "assessment_docx": assessment_docx_path, "zip": zip_path,
                    }
                    st.session_state.stage = "export_done"
                    st.rerun()
                except Exception as e:
                    st.error(f"An error occurred during export: {e}")

    with col2:
        if not ev.approved and st.session_state.revision_rounds < 2:
            if st.button("Run auto-revision", use_container_width=True):
                with st.spinner("Revising content based on evaluator feedback..."):
                    try:
                        revised = revision_agent(
                            st.session_state.plan, st.session_state.teaching_kit,
                            st.session_state.presentation, st.session_state.assessment,
                            ev, st.session_state.context
                        )
                        st.session_state.teaching_kit = revised.teaching_kit
                        st.session_state.presentation = revised.presentation
                        st.session_state.assessment = revised.assessment
                        st.session_state.revision_rounds += 1

                        new_eval = evaluator_agent(
                            st.session_state.plan, st.session_state.teaching_kit,
                            st.session_state.presentation, st.session_state.assessment,
                            st.session_state.context
                        )
                        st.session_state.evaluation = new_eval
                        st.rerun()
                    except Exception as e:
                        st.error(f"An error occurred during revision: {e}")

    with col3:
        if st.button("Manual edit (back to slides)", use_container_width=True):
            st.session_state.stage = "presentation_review"
            st.rerun()


# ============================================================
# STAGE: EXPORT DONE
# ============================================================
elif st.session_state.stage == "export_done":
    files = st.session_state.files
    plan = st.session_state.plan
    score = st.session_state.evaluation.score if st.session_state.evaluation else "—"
    theme_name = THEMES.get(st.session_state.theme_key, THEMES["slate"])["name"]

    st.markdown(
        """
        <div class='sage-success-hero'>
            <h2>✦ Teaching Kit Ready</h2>
            <p>Your teaching package has been successfully generated. Download any file individually, or grab the full ZIP package.</p>
        </div>
        """,
        unsafe_allow_html=True
    )

    c1, c2, c3 = st.columns(3)
    c1.metric("Topic", plan.topic)
    c2.metric("Final score", f"{score}/100")
    c3.metric("Theme", theme_name)

    st.write("")
    st.markdown("##### Files ready for download")

    d1, d2, d3, d4 = st.columns(4)
    with open(files["pptx"], "rb") as f:
        d1.download_button(
            "Presentation (PPTX)", f,
            file_name=os.path.basename(files["pptx"]),
            use_container_width=True
        )
    with open(files["kit_docx"], "rb") as f:
        d2.download_button(
            "Teaching Kit (DOCX)", f,
            file_name=os.path.basename(files["kit_docx"]),
            use_container_width=True
        )
    with open(files["assessment_docx"], "rb") as f:
        d3.download_button(
            "Assessment + Key (DOCX)", f,
            file_name=os.path.basename(files["assessment_docx"]),
            use_container_width=True
        )
    with open(files["zip"], "rb") as f:
        d4.download_button(
            "Full Package (ZIP)", f,
            file_name=os.path.basename(files["zip"]),
            use_container_width=True
        )

    st.divider()
    _s1, mid, _s2 = st.columns([1, 1, 1])
    with mid:
        if st.button("Generate a new kit", type="primary", use_container_width=True):
            for k, v in DEFAULTS.items():
                if k != "runs_this_session":
                    st.session_state[k] = v
            st.rerun()